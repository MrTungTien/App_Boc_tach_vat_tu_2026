import io
import math
import re
import unicodedata

import pandas as pd
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

st.set_page_config(page_title="Bóc tách & gộp vật tư cơ khí", layout="wide")
st.title("HỆ THỐNG TỰ ĐỘNG BÓC TÁCH & GỘP VẬT TƯ CƠ KHÍ")

CAY_MM = 6000  # chiều dài 1 cây thép hình (mm)
QC = "Quy cách/Ký mã hiệu"
COT_KQ = ["TT", "Tên gọi", QC, "Đơn vị tính", "Số lượng", "Đơn giá", "Thành tiền", "Ghi chú"]

# Từ khóa nhận diện tiêu đề cột (đã bỏ dấu, chữ thường)
TU_KHOA = {
    "nd": ["noi dung", "ten goi", "ten vat tu", "ten chi tiet", "dien giai", "ten"],
    "qc": ["quy cach", "ky ma hieu", "ma hieu", "ky hieu"],
    "dvt": ["dvt", "don vi"],
    "sl": ["so luong", "sl"],
    "qd": ["quy doi", "chieu dai"],
}
SO_DONG_TIM_TIEU_DE = 30


# ---------- Hàm hỗ trợ ----------
def bo_dau(s) -> str:
    s = unicodedata.normalize("NFD", str(s))
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return " ".join(s.replace("đ", "d").replace("Đ", "D").lower().split())


def chuan_hoa_text(x) -> str:
    """In hoa, xóa khoảng trắng thừa."""
    if pd.isna(x):
        return ""
    return " ".join(str(x).split()).upper()


def chuan_hoa_qc(x) -> str:
    """Quy cách: in hoa, bỏ khoảng trắng, đổi × và * thành X (V50 x 5 = V50X5)."""
    s = chuan_hoa_text(x).replace("×", "X").replace("*", "X")
    return re.sub(r"\s+", "", s)


def la_so(s: str) -> bool:
    return bool(re.fullmatch(r"[\d.,]+", s))


def tim_cot(o_tieu_de: list) -> dict:
    """Tìm vị trí các cột trong 1 dòng tiêu đề. Trả về {khóa: chỉ số cột}."""
    vt = {}
    for i, o in enumerate(o_tieu_de):
        k = bo_dau(o) if pd.notna(o) else ""
        if not k:
            continue
        for khoa in ("qd", "qc", "dvt", "sl", "nd"):  # thứ tự ưu tiên
            if khoa in vt:
                continue
            if khoa == "sl" and "quy doi" in k:
                continue
            if khoa == "nd" and any(t in k for t in ("don vi", "dvt", "don gia", "quy cach", "ma hieu")):
                continue
            if any(t == k or t in k.split("/") or (len(t) > 3 and t in k) or k.startswith(t + " ")
                   for t in TU_KHOA[khoa]):
                vt[khoa] = i
                break
    return vt


def doc_bang(raw: pd.DataFrame):
    """Tìm dòng tiêu đề trong sheet và trả về (dữ liệu chuẩn hóa, dòng tiêu đề) hoặc (None, lý do)."""
    for r in range(min(SO_DONG_TIM_TIEU_DE, len(raw))):
        vt = tim_cot(list(raw.iloc[r]))
        if all(k in vt for k in ("nd", "dvt", "sl")):
            dl = raw.iloc[r + 1:]
            d = pd.DataFrame({
                "Nội dung": dl.iloc[:, vt["nd"]].map(chuan_hoa_text),
                QC: dl.iloc[:, vt["qc"]].map(chuan_hoa_qc) if "qc" in vt else "",
                "ĐVT": dl.iloc[:, vt["dvt"]].map(chuan_hoa_text),
                "Số lượng": pd.to_numeric(dl.iloc[:, vt["sl"]], errors="coerce").fillna(0),
                "Quy đổi (mm)": pd.to_numeric(dl.iloc[:, vt["qd"]], errors="coerce") if "qd" in vt else float("nan"),
            })
            return d.reset_index(drop=True), int(raw.index[r]) + 1  # số dòng trong Excel
    return None, "Không tìm thấy dòng tiêu đề có đủ cột Tên gọi/Nội dung, ĐVT, Số lượng"


def nhan_he_so_cum(d: pd.DataFrame) -> pd.DataFrame:
    """Dòng ĐVT = CỤM đặt hệ số nhân cho các dòng bên dưới đến khi gặp Cụm mới (trong cùng sheet)."""
    he_so, hs = 1.0, []
    for dvt, sl in zip(d["ĐVT"], d["Số lượng"]):
        if dvt == "CỤM":
            he_so = sl
        hs.append(he_so)
    d = d.copy()
    d["Nhu cầu"] = d["Số lượng"] * pd.Series(hs, index=d.index)
    # Bỏ dòng Cụm, dòng trống, dòng không có ĐVT (tổng cộng, ghi chú...), dòng đánh số cột (1,2,3...)
    giu = (d["ĐVT"] != "CỤM") & (d["Nội dung"] != "") & (d["ĐVT"] != "") & ~d["ĐVT"].map(la_so)
    return d[giu]


def doc_tat_ca(files):
    """Đọc mọi file, mọi sheet. Trả về (bảng chi tiết gộp, nhật ký từng sheet)."""
    ds, nk = [], []
    for f in files:
        try:
            if f.name.lower().endswith(".csv"):
                sheets = None
                for enc in ("utf-8-sig", "utf-8", "cp1258", "latin-1"):
                    try:
                        f.seek(0)
                        sheets = {"(csv)": pd.read_csv(f, header=None, encoding=enc, dtype=object)}
                        break
                    except UnicodeDecodeError:
                        continue
                if sheets is None:
                    raise ValueError("Không đọc được mã hóa CSV")
            else:
                f.seek(0)
                sheets = pd.read_excel(f, sheet_name=None, header=None, engine="openpyxl", dtype=object)
        except Exception as e:
            nk.append({"File": f.name, "Sheet": "", "Dòng tiêu đề": "", "Số dòng vật tư": 0,
                       "Trạng thái": f"Lỗi đọc file: {e}"})
            continue

        for ten_sheet, raw in sheets.items():
            raw = raw.dropna(how="all").dropna(axis=1, how="all")
            if raw.empty:
                nk.append({"File": f.name, "Sheet": ten_sheet, "Dòng tiêu đề": "", "Số dòng vật tư": 0,
                           "Trạng thái": "Bỏ qua: sheet trống"})
                continue
            d, ttd = doc_bang(raw)
            if d is None:
                nk.append({"File": f.name, "Sheet": ten_sheet, "Dòng tiêu đề": "", "Số dòng vật tư": 0,
                           "Trạng thái": f"Bỏ qua: {ttd}"})
                continue
            ct = nhan_he_so_cum(d)
            ct = ct.assign(File=f.name, Sheet=ten_sheet)
            ds.append(ct)
            nk.append({"File": f.name, "Sheet": ten_sheet, "Dòng tiêu đề": ttd, "Số dòng vật tư": len(ct),
                       "Trạng thái": "Đã đọc"})
    chi_tiet = pd.concat(ds, ignore_index=True) if ds else pd.DataFrame()
    return chi_tiet, pd.DataFrame(nk)


def tong_hop(ct: pd.DataFrame) -> pd.DataFrame:
    """Gom theo Tên gọi + Quy cách/Ký mã hiệu. Thép hình (ĐVT thanh) quy ra cây 6m làm tròn lên."""
    is_thanh = ct["ĐVT"] == "THANH"

    hinh = ct[is_thanh].copy()
    hinh["Dài (mm)"] = hinh["Quy đổi (mm)"].fillna(0) * hinh["Nhu cầu"]
    hinh = hinh.groupby(["Nội dung", QC], as_index=False)["Dài (mm)"].sum()
    hinh["ĐVT"] = "CÂY"
    hinh["Số lượng"] = hinh["Dài (mm)"].map(lambda v: math.ceil(v / CAY_MM - 1e-9))

    khac = (
        ct[~is_thanh]
        .groupby(["Nội dung", QC, "ĐVT"], as_index=False)["Nhu cầu"]
        .sum()
        .rename(columns={"Nhu cầu": "Số lượng"})
    )

    bang = pd.concat(
        [hinh[["Nội dung", QC, "ĐVT", "Số lượng"]], khac[["Nội dung", QC, "ĐVT", "Số lượng"]]],
        ignore_index=True,
    ).sort_values(["Nội dung", QC], kind="stable").reset_index(drop=True)
    bang.insert(0, "TT", range(1, len(bang) + 1))
    bang["Đơn giá"] = float("nan")
    bang["Thành tiền"] = float("nan")
    bang["Ghi chú"] = ""
    bang = bang.rename(columns={"Nội dung": "Tên gọi", "ĐVT": "Đơn vị tính"})
    return bang[COT_KQ]


def tinh_thanh_tien(bang: pd.DataFrame) -> pd.DataFrame:
    b = bang.copy()
    b["Đơn giá"] = pd.to_numeric(b["Đơn giá"], errors="coerce")
    b["Số lượng"] = pd.to_numeric(b["Số lượng"], errors="coerce")
    b["Thành tiền"] = b["Số lượng"] * b["Đơn giá"]
    b["Ghi chú"] = b["Ghi chú"].fillna("")
    b[QC] = b[QC].fillna("")
    return b


def xuat_excel(bang: pd.DataFrame, nhat_ky: pd.DataFrame) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Bang_tong_hop"
    vien = Border(*(Side(style="thin"),) * 4)
    thuong = Font(name="Times New Roman", size=12)
    dam = Font(name="Times New Roman", size=12, bold=True)

    ws.append(COT_KQ)
    ws.row_dimensions[1].height = 48
    for c in ws[1]:
        c.font = dam
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for _, r in bang.iterrows():
        dong = ws.max_row + 1
        don_gia = None if pd.isna(r["Đơn giá"]) else float(r["Đơn giá"])
        ws.append([
            int(r["TT"]), r["Tên gọi"], r[QC] or None, r["Đơn vị tính"], float(r["Số lượng"]), don_gia,
            f'=IF(F{dong}="","",E{dong}*F{dong})',  # Thành tiền tự tính khi nhập đơn giá
            r["Ghi chú"] or None,
        ])

    cuoi = ws.max_row
    ws.append([None, "TỔNG CỘNG", None, None, None, None, f"=SUM(G2:G{cuoi})", None])
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=len(COT_KQ)):
        for c in row:
            c.border = vien
            if c.row > 1:
                c.font = dam if c.row == ws.max_row else thuong
                c.alignment = Alignment(vertical="center", wrap_text=True)
            if c.column in (5, 6, 7) and c.row > 1:
                c.number_format = "#,##0.##" if c.column == 5 else "#,##0"
            if c.column in (1, 4) and c.row > 1:
                c.alignment = Alignment(horizontal="center", vertical="center")

    for i, w in enumerate([8, 30, 20, 14, 12, 16, 18, 18], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    for i in range(2, ws.max_row + 1):
        ws.row_dimensions[i].height = 22
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"

    # Sheet phụ: danh sách file/sheet đã đọc để đối chiếu
    ws2 = wb.create_sheet("Nguon_du_lieu")
    ws2.append(list(nhat_ky.columns))
    for r in nhat_ky.itertuples(index=False):
        ws2.append(list(r))
    for c in ws2[1]:
        c.font = dam
    for i, w in enumerate([40, 25, 14, 16, 60], start=1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------- Giao diện ----------
files = st.file_uploader(
    "Kéo thả một hoặc nhiều file Excel (.xlsx) / CSV. Mỗi file có thể có nhiều sheet.",
    type=["xlsx", "xlsm", "csv"],
    accept_multiple_files=True,
)

if files:
    st.write(f"Đã chọn **{len(files)}** file.")
    if st.button("Xử lý & Gom nhóm vật tư", type="primary"):
        with st.spinner("Đang đọc các file..."):
            ct, nk = doc_tat_ca(files)
        st.session_state["nk"] = nk
        st.session_state["bang"] = tong_hop(ct) if not ct.empty else None
        st.session_state["lan"] = st.session_state.get("lan", 0) + 1  # làm mới bảng nhập

    if "nk" in st.session_state:
        nk = st.session_state["nk"]
        doc_duoc = (nk["Trạng thái"] == "Đã đọc").sum()
        bo_qua = len(nk) - doc_duoc
        st.info(f"Đọc được {doc_duoc} sheet, bỏ qua {bo_qua} sheet.")
        with st.expander("Chi tiết từng file / sheet", expanded=bo_qua > 0):
            st.dataframe(nk, hide_index=True, use_container_width=True)

    if st.session_state.get("bang") is not None:
        st.subheader("Bảng tổng hợp vật tư")
        st.caption("Gom theo Tên gọi + Quy cách/Ký mã hiệu. Thép hình tính theo cây 6m (làm tròn lên). "
                   "Nhập Đơn giá và Ghi chú trực tiếp vào bảng.")
        sua = st.data_editor(
            st.session_state["bang"],
            key=f"editor_{st.session_state['lan']}",
            hide_index=True,
            use_container_width=True,
            disabled=["TT", "Tên gọi", "Đơn vị tính", "Số lượng", "Thành tiền"],
            column_config={
                "Đơn giá": st.column_config.NumberColumn(min_value=0, step=1, format="%.0f"),
                "Thành tiền": st.column_config.NumberColumn(format="%.0f"),
                "Ghi chú": st.column_config.TextColumn(),
                QC: st.column_config.TextColumn(),
            },
        )
        kq = tinh_thanh_tien(sua)
        st.metric("Tổng cộng thành tiền", f"{kq['Thành tiền'].sum(skipna=True):,.0f}")

        st.download_button(
            "Tải về file Excel tổng hợp",
            data=xuat_excel(kq, st.session_state["nk"]),
            file_name="bang_tong_hop_vat_tu.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    elif "nk" in st.session_state:
        st.error("Không có sheet nào đọc được. Xem bảng chi tiết ở trên để biết lý do.")
else:
    st.info("Hãy tải file dữ liệu lên để bắt đầu.")
